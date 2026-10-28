import shutil
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from app.core.errors import ApiError
from app.modules.ingestion.extraction import (
    OcrOptions,
    detect_format,
    detect_pdf_type,
    extract_document,
)

PARAGRAPH = (
    "Le programme d'adaptation agricole en Ethiopie finance l'irrigation, "
    "la gestion durable de l'eau et la resilience des exploitations familiales."
)


def _pdf(pages: list[str]) -> bytes:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    return document.tobytes()


def _scanned_pdf(text: str) -> bytes:
    """Une page de texte rendue en image, sans couche texte."""
    source = pymupdf.open(stream=_pdf([text]), filetype="pdf")
    pixmap = source[0].get_pixmap(dpi=200)
    scanned = pymupdf.open()
    page = scanned.new_page(width=source[0].rect.width, height=source[0].rect.height)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    return scanned.tobytes()


def test_detects_format_from_content_not_extension() -> None:
    assert detect_format(_pdf(["x"]), "rapport.docx") == "pdf"
    with pytest.raises(ApiError) as error:
        detect_format(b"plain text", "notes.pdf")
    assert error.value.code == "UNSUPPORTED_FORMAT"


def test_native_pdf_keeps_text_per_page_with_offsets() -> None:
    extracted = extract_document(_pdf([PARAGRAPH, PARAGRAPH * 2]), "prodoc.pdf")
    assert extracted.format == "pdf"
    assert extracted.pdf_type == "natif"
    assert extracted.total_pages == 2
    assert "irrigation" in extracted.pages[0].text
    assert not extracted.pages[0].is_ocr
    assert extracted.pages[1].char_start == len(extracted.pages[0].text) + 2
    assert extracted.exploitable_pages(100) == 2


def _table_page(document: pymupdf.Document, rows: list[tuple[str, str]]) -> None:
    """Un paragraphe suivi d'un tableau tracé (filets vectoriels)."""
    page = document.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 300), PARAGRAPH, fontsize=11)
    top, height, columns = 320, 24, (50, 300, 550)
    for index in range(len(rows) + 1):
        y = top + index * height
        page.draw_line((columns[0], y), (columns[-1], y))
    for x in columns:
        page.draw_line((x, top), (x, top + len(rows) * height))
    for index, row in enumerate(rows):
        for x, value in zip(columns, row, strict=False):
            page.insert_text((x + 5, top + index * height + 16), value, fontsize=10)


def test_pdf_tables_are_only_searched_on_pages_with_drawings(monkeypatch) -> None:
    document = pymupdf.open(stream=_pdf([PARAGRAPH * 2]), filetype="pdf")
    _table_page(document, [("Activite", "Budget"), ("Irrigation", "120 000")])
    calls: list[int] = []
    find_tables = pymupdf.Page.find_tables

    def counting(page, *args, **kwargs):
        calls.append(page.number)
        return find_tables(page, *args, **kwargs)

    monkeypatch.setattr(pymupdf.Page, "find_tables", counting)
    extracted = extract_document(document.tobytes(), "budget.pdf")

    assert calls == [1]
    assert "Irrigation | 120 000" in extracted.pages[1].text


def test_pdf_type_follows_native_text_and_images() -> None:
    text = "x" * 300
    assert detect_pdf_type([text, text], [False, True]) == "natif"
    assert detect_pdf_type([text, ""], [False, True]) == "mixte"
    assert detect_pdf_type(["", "court"], [True, True]) == "scan"
    assert detect_pdf_type([], []) == "scan"


def test_empty_pages_are_not_exploitable() -> None:
    extracted = extract_document(_pdf([PARAGRAPH, ""]), "partiel.pdf")
    assert extracted.total_pages == 2
    assert extracted.exploitable_pages(100) == 1


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract absent")
def test_scanned_pdf_goes_through_ocr() -> None:
    extracted = extract_document(
        _scanned_pdf(PARAGRAPH * 2), "cadre-logique-scanne.pdf", OcrOptions(dpi=200)
    )
    assert extracted.pdf_type == "scan"
    page = extracted.pages[0]
    assert page.is_ocr
    assert not page.ocr_failed
    assert "irrigation" in page.text.lower()


def test_password_protected_pdf_is_rejected_explicitly() -> None:
    document = pymupdf.open(stream=_pdf([PARAGRAPH]), filetype="pdf")
    protected = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner"
    )
    with pytest.raises(ApiError) as error:
        extract_document(protected, "protege.pdf")
    assert error.value.code == "PASSWORD_PROTECTED_PDF"


def test_corrupted_file_is_reported() -> None:
    with pytest.raises(ApiError) as error:
        extract_document(b"%PDF-1.7 broken content", "casse.pdf")
    assert error.value.code == "CORRUPTED_FILE"


def test_docx_keeps_headings_lists_and_tables_in_order() -> None:
    document = Document()
    document.add_heading("1. Contexte", level=1)
    document.add_paragraph(PARAGRAPH)
    document.add_paragraph("Renforcer les capacités", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Activité"
    table.cell(0, 1).text = "Budget"
    table.cell(1, 0).text = "Irrigation"
    table.cell(1, 1).text = "120 000"
    buffer = BytesIO()
    document.save(buffer)

    extracted = extract_document(buffer.getvalue(), "prodoc.docx")
    text = extracted.pages[0].text
    assert extracted.format == "docx"
    assert extracted.pages[0].section_title == "1. Contexte"
    assert text.index("1. Contexte") < text.index("irrigation")
    assert "- Renforcer les capacités" in text
    assert "Irrigation | 120 000" in text


def test_xlsx_gives_one_page_per_sheet_with_cell_positions() -> None:
    workbook = Workbook()
    budget = workbook.active
    budget.title = "Budget"
    budget["A1"] = "Activité"
    budget["B1"] = "Montant"
    budget["A2"] = "Irrigation"
    budget["B2"] = 120000
    indicators = workbook.create_sheet("Indicateurs")
    indicators["A1"] = "Ménages"
    buffer = BytesIO()
    workbook.save(buffer)

    extracted = extract_document(buffer.getvalue(), "budget.xlsx")
    assert extracted.format == "xlsx"
    assert [page.section_title for page in extracted.pages] == ["Budget", "Indicateurs"]
    assert "A2=Irrigation | B2=120000" in extracted.pages[0].text
