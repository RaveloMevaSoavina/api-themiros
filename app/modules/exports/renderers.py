"""Mise en page du rapport d'un run : PDF, Word, XLSX et CSV (US-22.1).

Chaque format reprend le même `Report`, donc les mêmes sections, limites et
notes : seul le rendu change. Le PDF est un rapport mis en forme, le Word un
document retravaillable, le XLSX un onglet par écran et le CSV un tableau
unique, lisible par un tableur francophone.
"""

import csv
import io
import re
from collections.abc import Callable
from html import escape

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.exports.report import Note, PillarScore, Report

WARNING = "B42318"
MUTED = "667085"


def paragraphs(text: str | None) -> list[str]:
    """Conclusion narrative : un paragraphe par bloc (prompt v1.1)."""
    return [
        block.strip() for block in re.split(r"\n\s*\n", text or "") if block.strip()
    ]


def filename(report: Report, extension: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", report.workspace.lower()).strip("-")[:40]
    date = report.date(report.run_date).replace("/", "-")
    return f"themiros-{slug or 'rapport'}-{date}.{extension}"


# --- PDF ---------------------------------------------------------------------

# Pas de background-color : le moteur HTML de MuPDF le répète en bandes
# parasites sur les éléments qui suivent un tableau.
PDF_CSS = f"""
* {{ font-family: sans-serif; }}
body {{ font-size: 9.5pt; line-height: 1.35; color: #101828; }}
h1 {{ font-size: 20pt; margin: 0 0 2pt 0; }}
h2 {{ font-size: 13pt; margin: 16pt 0 4pt 0; color: #101828; }}
h3 {{ font-size: 11pt; margin: 10pt 0 3pt 0; }}
p {{ margin: 0 0 4pt 0; }}
.sub {{ color: #{MUTED}; font-size: 11pt; }}
.muted {{ color: #{MUTED}; font-size: 8.5pt; }}
.banner {{ color: #{WARNING}; font-size: 10pt; margin: 8pt 0; }}
.refs {{ color: #{MUTED}; font-size: 8pt; }}
table {{ border-collapse: collapse; width: 100%; margin: 2pt 0 6pt 0; }}
th {{ text-align: left; font-size: 8.5pt; font-weight: bold;
      border: 0.5pt solid #D0D5DD; padding: 2pt 3pt; }}
td {{ font-size: 8.5pt; border: 0.5pt solid #D0D5DD; padding: 2pt 3pt;
      vertical-align: top; }}
li {{ margin-bottom: 3pt; }}
"""


def _e(value: object) -> str:
    return escape(str(value))


def _table(headers: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{_e(header)}</th>" for header in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{_e(cell)}</td>" for cell in row) + "</tr>"
        for row in rows
    )
    return f"<table><tr>{head}</tr>{body}</table>"


def _refs(report: Report, numbers: list[int]) -> str:
    label = report.labels["columns"]["notes"]
    return f"<p class='refs'>{_e(label)} {_e(report.refs(numbers))}</p>"


def _origin(report: Report, origin: str, numbers: list[int]) -> str:
    refs = f" · {report.refs(numbers)}" if numbers else ""
    return f"<span class='refs'>{_e(origin + refs)}</span>"


def _html(report: Report) -> str:
    t = report.labels
    c = t["columns"]
    sections = t["sections"]
    html: list[str] = [
        f"<h1>{_e(t['title'])}</h1>",
        f"<p class='sub'>{_e(report.workspace)}</p>",
    ]
    if report.exploratory:
        html.append(
            f"<div class='banner'><b>{_e(t['exploratory']['title'])}</b><br/>"
            f"{_e(report.exploratory_text)}</div>"
        )
    html.append(
        "<table>"
        + "".join(
            f"<tr><th>{_e(name)}</th><td>{_e(value)}</td></tr>"
            for name, value in report.meta_rows()
        )
        + "</table>"
    )

    html.append(f"<h2>{_e(sections['limits'])}</h2><ul>")
    html.extend(f"<li>{_e(line)}</li>" for line in report.limits)
    html.append("</ul>")

    score_headers = [c["score"], c["evolution"], c["confidence"], c["documents"]]

    html.append(f"<h2>{_e(sections['overview'])}</h2>")
    html.append(f"<p class='muted'>{_e(t['layers']['a'])}</p>")
    if report.pillars:
        html.append(
            _table(
                [c["pillar"], *score_headers],
                [[p.name, *_score_cells(report, p.latest)] for p in report.pillars],
            )
        )
        average = report.number(report.average_score)
        html.append(f"<p><b>{_e(report.labeled(t['average'], average))}</b></p>")
    else:
        html.append(f"<p>{_e(t['no_scores'])}</p>")

    html.append(f"<h2>{_e(sections['pillars'])}</h2>")
    if not report.pillars:
        html.append(f"<p>{_e(t['no_scores'])}</p>")
    for pillar in report.pillars:
        html.append(f"<h3>{_e(pillar.name)}</h3>")
        if pillar.description:
            html.append(f"<p class='muted'>{_e(pillar.description)}</p>")
        if report.multiple_versions:
            html.append(
                _table(
                    [c["version"], *score_headers],
                    [[s.version, *_score_cells(report, s)] for s in pillar.scores],
                )
            )
        else:
            score = pillar.latest
            summary = " · ".join(
                [
                    report.labeled(c["score"], report.score(score.score)),
                    report.labeled(c["confidence"], report.percent(score.confidence)),
                    report.labeled(c["documents"], score.documents),
                ]
            )
            html.append(f"<p>{_e(summary)}</p>")
        if pillar.variables:
            headers = [c["variable"]]
            if report.multiple_versions:
                headers.append(c["version"])
            headers += [c["state"], c["documents"], c["confidence"], c["notes"]]
            rows = []
            for variable in pillar.variables:
                row = [variable.label]
                if report.multiple_versions:
                    row.append(variable.version)
                row += [
                    t["state"].get(variable.state, variable.state),
                    str(variable.documents),
                    report.percent(variable.confidence),
                    report.refs(variable.notes),
                ]
                rows.append(row)
            html.append(_table(headers, rows))
        else:
            html.append(f"<p>{_e(t['no_variables'])}</p>")
        if pillar.notes:
            html.append(_refs(report, pillar.notes))

    html.append(f"<h2>{_e(sections['criteria'])}</h2>")
    html.append(f"<p class='muted'>{_e(t['layers']['b'])}</p>")
    if report.criteria:
        html.append(
            _table(
                [
                    c["code"],
                    c["criterion"],
                    c["note"],
                    c["status"],
                    c["sub_questions"],
                    c["ambiguous"],
                ],
                [
                    [
                        criterion.code,
                        criterion.name,
                        report.number(criterion.note, 1),
                        t["status"].get(criterion.status, criterion.status),
                        f"{criterion.documented}/{criterion.total}",
                        t["yes"] if criterion.ambiguous else t["no"],
                    ]
                    for criterion in report.criteria
                ],
            )
        )
        for criterion in report.criteria:
            html.append(f"<h3>{_e(criterion.code)} · {_e(criterion.name)}</h3>")
            html.append(f"<p>{_e(criterion.justification or t['empty'])}</p>")
            if criterion.notes:
                html.append(_refs(report, criterion.notes))
    else:
        html.append(f"<p>{_e(t['no_criteria'])}</p>")

    html.append(f"<h2>{_e(sections['alerts'])}</h2>")
    html.append(f"<p class='muted'>{_e(t['layers']['c'])}</p>")
    if not report.alerts:
        html.append(f"<p>{_e(t['no_alerts'])}</p>")
    for alert in report.alerts:
        severity = t["severity"].get(alert.severity, alert.severity)
        html.append(f"<h3>{_e(f'{severity} · {alert.rule} · {alert.scope}')}</h3>")
        html.append(f"<p>{_e(alert.message)}</p>")
        details = " · ".join(
            [
                report.labeled(
                    c["status"], t["alert_status"].get(alert.status, alert.status)
                ),
                report.labeled(c["qualified"], t["yes" if alert.qualified else "no"]),
            ]
        )
        html.append(f"<p class='muted'>{_e(details)}</p>")
        if alert.comment:
            html.append(
                f"<p><i>{_e(report.labeled(c['comment'], alert.comment))}</i></p>"
            )
        if alert.notes:
            html.append(_refs(report, alert.notes))

    html.append(f"<h2>{_e(sections['synthesis'])}</h2>")
    if report.layers.get("restitution"):
        for section, title in t["synthesis_sections"].items():
            html.append(f"<h3>{_e(title)}</h3>")
            items = [item for item in report.synthesis if item.section == section]
            if not items:
                empty = t["synthesis_empty"][section]
                html.append(f"<p class='muted'>{_e(empty)}</p>")
                continue
            html.append("<ol>")
            html.extend(
                f"<li>{_e(item.text)}<br/>"
                f"{_origin(report, item.origin, item.notes)}</li>"
                for item in items
            )
            html.append("</ol>")
        html.append(f"<h3>{_e(t['conclusion'])}</h3>")
        blocks = paragraphs(report.conclusion)
        html.extend(f"<p>{_e(block)}</p>" for block in blocks)
        if not blocks:
            html.append(f"<p class='muted'>{_e(t['no_conclusion'])}</p>")
    else:
        html.append(f"<p>{_e(t['limits']['restitution_missing'])}</p>")

    persona = t["personas"].get(report.persona, report.persona)
    html.append(f"<h2>{_e(sections['actions'].format(persona=persona))}</h2>")
    if report.actions:
        html.append("<ol>")
        html.extend(
            f"<li><b>{_e(action.text)}</b><br/>"
            f"{_e(t['because'].format(reason=action.justification))}<br/>"
            f"{_origin(report, action.origin, action.notes)}</li>"
            for action in report.actions
        )
        html.append("</ol>")
    else:
        html.append(f"<p>{_e(report.actions_empty())}</p>")

    html.append(f"<h2>{_e(sections['notes'])}</h2>")
    html.append(f"<p class='muted'>{_e(report.notes_intro)}</p>")
    if not report.notes:
        html.append(f"<p>{_e(t['no_notes'])}</p>")
    for note in report.notes:
        lines = [
            f"<b>[{note.number}]</b> {_e(_source(report, note))}",
            f"<span class='muted'>{_e(note.subject)}</span>",
            *(_e(report.quote(excerpt)) for excerpt in note.excerpts),
        ]
        omitted = report.omitted(note)
        if omitted:
            lines.append(f"<span class='muted'>{_e(omitted)}</span>")
        html.append(f"<p>{'<br/>'.join(lines)}</p>")
    return "".join(html)


def _evolution(report: Report, evolution: str | None, delta: float | None) -> str:
    if evolution is None:
        return report.labels["empty"]
    label = report.labels["evolution"].get(evolution, evolution)
    return f"{label} ({report.signed(delta)})" if delta is not None else label


def _score_cells(report: Report, score: PillarScore) -> list[str]:
    return [
        report.score(score.score),
        _evolution(report, score.evolution, score.delta),
        report.percent(score.confidence),
        str(score.documents),
    ]


def _source(report: Report, note: Note) -> str:
    parts = [note.document or report.labels["empty"]]
    if note.page is not None:
        parts.append(report.labels["page"].format(page=note.page))
    if note.section:
        parts.append(note.section)
    return " · ".join(parts)


def _footer(report: Report) -> str:
    parts = [report.labels["generator"], report.workspace, f"Run {report.run_id[:8]}"]
    if report.exploratory:
        parts.append(report.labels["exploratory"]["footer"])
    return " · ".join(parts)


def render_pdf(report: Report) -> bytes:
    story = pymupdf.Story(html=_html(report), user_css=PDF_CSS)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (48, 48, -48, -56)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()

    document = pymupdf.open("pdf", buffer.getvalue())
    footer = _footer(report)
    total = document.page_count
    for index, page in enumerate(document, start=1):
        bottom = page.rect.height - 28
        page.insert_text(
            (48, bottom),
            footer,
            fontsize=7.5,
            fontname="helv",
            color=(0.4, 0.44, 0.52),
        )
        number = report.labels["page_footer"].format(page=index, total=total)
        width = pymupdf.get_text_length(number, fontname="helv", fontsize=7.5)
        page.insert_text(
            (page.rect.width - 48 - width, bottom),
            number,
            fontsize=7.5,
            fontname="helv",
            color=(0.4, 0.44, 0.52),
        )
    document.set_metadata(
        {
            "title": f"{report.labels['title']} · {report.workspace}",
            "author": report.exported_by,
            "creator": report.labels["generator"],
            "producer": report.labels["generator"],
        }
    )
    return document.tobytes(garbage=3, deflate=True)


# --- Word --------------------------------------------------------------------


def _shade(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _docx_table(document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(8.5)
        _shade(cell, "F2F4F7")
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = ""
            cell.paragraphs[0].add_run(value).font.size = Pt(8.5)
    document.add_paragraph()


def _muted(document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    return paragraph


def _page_field(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def render_docx(report: Report) -> bytes:
    t = report.labels
    c = t["columns"]
    sections = t["sections"]
    document = Document()
    document.styles["Normal"].font.size = Pt(10)
    document.core_properties.title = f"{t['title']} · {report.workspace}"
    document.core_properties.author = report.exported_by

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = f"{_footer(report)} · "
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_field(footer)
    for run in footer.runs:
        run.font.size = Pt(7.5)

    document.add_heading(t["title"], 0)
    document.add_paragraph(report.workspace)
    if report.exploratory:
        banner = document.add_table(rows=1, cols=1).rows[0].cells[0]
        _shade(banner, "FEF3F2")
        banner.text = ""
        title = banner.paragraphs[0].add_run(t["exploratory"]["title"])
        title.bold = True
        title.font.color.rgb = RGBColor.from_string(WARNING)
        text = banner.add_paragraph().add_run(report.exploratory_text)
        text.font.color.rgb = RGBColor.from_string(WARNING)
        document.add_paragraph()
    _docx_table(
        document,
        [c["field"], c["value"]],
        [[name, value] for name, value in report.meta_rows()],
    )

    document.add_heading(sections["limits"], 1)
    for line in report.limits:
        document.add_paragraph(line, style="List Bullet")

    score_headers = [c["score"], c["evolution"], c["confidence"], c["documents"]]
    document.add_heading(sections["overview"], 1)
    _muted(document, t["layers"]["a"])
    if report.pillars:
        _docx_table(
            document,
            [c["pillar"], *score_headers],
            [[p.name, *_score_cells(report, p.latest)] for p in report.pillars],
        )
        average = report.labeled(t["average"], report.number(report.average_score))
        document.add_paragraph().add_run(average).bold = True
    else:
        document.add_paragraph(t["no_scores"])

    document.add_heading(sections["pillars"], 1)
    if not report.pillars:
        document.add_paragraph(t["no_scores"])
    for pillar in report.pillars:
        document.add_heading(pillar.name, 2)
        if pillar.description:
            _muted(document, pillar.description)
        _docx_table(
            document,
            [c["version"], *score_headers],
            [[s.version, *_score_cells(report, s)] for s in pillar.scores],
        )
        if pillar.variables:
            headers = [c["variable"]]
            if report.multiple_versions:
                headers.append(c["version"])
            headers += [c["state"], c["documents"], c["confidence"], c["notes"]]
            rows = []
            for variable in pillar.variables:
                row = [variable.label]
                if report.multiple_versions:
                    row.append(variable.version)
                row += [
                    t["state"].get(variable.state, variable.state),
                    str(variable.documents),
                    report.percent(variable.confidence),
                    report.refs(variable.notes),
                ]
                rows.append(row)
            _docx_table(document, headers, rows)
        else:
            document.add_paragraph(t["no_variables"])
        if pillar.notes:
            _muted(document, f"{c['notes']} {report.refs(pillar.notes)}")

    document.add_heading(sections["criteria"], 1)
    _muted(document, t["layers"]["b"])
    if report.criteria:
        _docx_table(
            document,
            [
                c["code"],
                c["criterion"],
                c["note"],
                c["status"],
                c["sub_questions"],
                c["ambiguous"],
            ],
            [
                [
                    criterion.code,
                    criterion.name,
                    report.number(criterion.note, 1),
                    t["status"].get(criterion.status, criterion.status),
                    f"{criterion.documented}/{criterion.total}",
                    t["yes"] if criterion.ambiguous else t["no"],
                ]
                for criterion in report.criteria
            ],
        )
        for criterion in report.criteria:
            document.add_heading(f"{criterion.code} · {criterion.name}", 3)
            document.add_paragraph(criterion.justification or t["empty"])
            if criterion.notes:
                _muted(document, f"{c['notes']} {report.refs(criterion.notes)}")
    else:
        document.add_paragraph(t["no_criteria"])

    document.add_heading(sections["alerts"], 1)
    _muted(document, t["layers"]["c"])
    if not report.alerts:
        document.add_paragraph(t["no_alerts"])
    for alert in report.alerts:
        severity = t["severity"].get(alert.severity, alert.severity)
        document.add_heading(f"{severity} · {alert.rule} · {alert.scope}", 3)
        document.add_paragraph(alert.message)
        status = t["alert_status"].get(alert.status, alert.status)
        _muted(
            document,
            f"{report.labeled(c['status'], status)} · "
            f"{report.labeled(c['qualified'], t['yes' if alert.qualified else 'no'])}",
        )
        if alert.comment:
            document.add_paragraph().add_run(
                report.labeled(c["comment"], alert.comment)
            ).italic = True
        if alert.notes:
            _muted(document, f"{c['notes']} {report.refs(alert.notes)}")

    document.add_heading(sections["synthesis"], 1)
    if report.layers.get("restitution"):
        for section, title in t["synthesis_sections"].items():
            document.add_heading(title, 2)
            items = [item for item in report.synthesis if item.section == section]
            if not items:
                _muted(document, t["synthesis_empty"][section])
            for item in items:
                document.add_paragraph(item.text, style="List Number")
                _muted(
                    document,
                    item.origin
                    + (f" · {report.refs(item.notes)}" if item.notes else ""),
                )
        document.add_heading(t["conclusion"], 2)
        blocks = paragraphs(report.conclusion)
        for block in blocks:
            document.add_paragraph(block)
        if not blocks:
            _muted(document, t["no_conclusion"])
    else:
        document.add_paragraph(t["limits"]["restitution_missing"])

    persona = t["personas"].get(report.persona, report.persona)
    document.add_heading(sections["actions"].format(persona=persona), 1)
    if not report.actions:
        document.add_paragraph(report.actions_empty())
    for action in report.actions:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(action.text).bold = True
        document.add_paragraph(t["because"].format(reason=action.justification))
        _muted(
            document,
            action.origin + (f" · {report.refs(action.notes)}" if action.notes else ""),
        )

    document.add_heading(sections["notes"], 1)
    _muted(document, report.notes_intro)
    if not report.notes:
        document.add_paragraph(t["no_notes"])
    for note in report.notes:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[{note.number}] ").bold = True
        paragraph.add_run(_source(report, note))
        subject = paragraph.add_run(f"\n{note.subject}")
        subject.font.size = Pt(8.5)
        subject.font.color.rgb = RGBColor.from_string(MUTED)
        for excerpt in note.excerpts:
            paragraph.add_run(f"\n{report.quote(excerpt)}").font.size = Pt(9)
        omitted = report.omitted(note)
        if omitted:
            run = paragraph.add_run(f"\n{omitted}")
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(MUTED)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Tableaux ----------------------------------------------------------------

Cell = str | int | float | None


def _sheets(report: Report) -> list[tuple[str, list[str], list[list[Cell]]]]:
    """Un tableau par écran, partagé par le XLSX (onglets)."""
    t = report.labels
    c = t["columns"]
    sheets = t["sheets"]
    persona = t["personas"].get(report.persona, report.persona)
    conclusion = [
        [t["conclusion"], index, block, None, None]
        for index, block in enumerate(paragraphs(report.conclusion), start=1)
    ]
    return [
        (
            sheets["overview"],
            [
                c["pillar"],
                c["version"],
                c["score"],
                c["delta"],
                c["evolution"],
                c["confidence"],
                c["documents"],
                c["notes"],
            ],
            [
                [
                    pillar.name,
                    score.version,
                    score.score,
                    score.delta,
                    t["evolution"].get(score.evolution or "", None),
                    score.confidence,
                    score.documents,
                    report.refs(pillar.notes),
                ]
                for pillar in report.pillars
                for score in pillar.scores
            ],
        ),
        (
            sheets["variables"],
            [
                c["pillar"],
                c["code"],
                c["variable"],
                c["version"],
                c["state"],
                c["value"],
                c["documents"],
                c["confidence"],
                c["notes"],
            ],
            [
                [
                    pillar.name,
                    variable.code,
                    variable.label,
                    variable.version,
                    t["state"].get(variable.state, variable.state),
                    variable.value,
                    variable.documents,
                    variable.confidence,
                    report.refs(variable.notes),
                ]
                for pillar in report.pillars
                for variable in pillar.variables
            ],
        ),
        (
            sheets["criteria"],
            [
                c["code"],
                c["criterion"],
                c["note"],
                c["status"],
                c["sub_questions"],
                c["ambiguous"],
                c["justification"],
                c["notes"],
            ],
            [
                [
                    criterion.code,
                    criterion.name,
                    criterion.note,
                    t["status"].get(criterion.status, criterion.status),
                    f"{criterion.documented}/{criterion.total}",
                    t["yes"] if criterion.ambiguous else t["no"],
                    criterion.justification,
                    report.refs(criterion.notes),
                ]
                for criterion in report.criteria
            ],
        ),
        (
            sheets["alerts"],
            [
                c["severity"],
                c["rule"],
                c["type"],
                c["scope"],
                c["status"],
                c["qualified"],
                c["message"],
                c["comment"],
                c["notes"],
            ],
            [
                [
                    t["severity"].get(alert.severity, alert.severity),
                    alert.rule,
                    alert.type,
                    alert.scope,
                    t["alert_status"].get(alert.status, alert.status),
                    t["yes"] if alert.qualified else t["no"],
                    alert.message,
                    alert.comment,
                    report.refs(alert.notes),
                ]
                for alert in report.alerts
            ],
        ),
        (
            sheets["synthesis"],
            [c["section"], c["position"], c["text"], c["origin"], c["notes"]],
            [
                [
                    t["synthesis_sections"].get(item.section, item.section),
                    item.position,
                    item.text,
                    item.origin,
                    report.refs(item.notes),
                ]
                for item in report.synthesis
            ]
            + conclusion,
        ),
        (
            sheets["actions"],
            [
                c["persona"],
                c["position"],
                c["action"],
                c["reason"],
                c["origin"],
                c["notes"],
            ],
            [
                [
                    persona,
                    action.position,
                    action.text,
                    action.justification,
                    action.origin,
                    report.refs(action.notes),
                ]
                for action in report.actions
            ],
        ),
        (
            sheets["notes"],
            [
                c["number"],
                c["document"],
                c["page"],
                c["section"],
                c["subject"],
                c["excerpt"],
            ],
            [
                [
                    note.number,
                    note.document,
                    note.page,
                    note.section,
                    note.subject,
                    note.excerpt,
                ]
                for note in report.notes
            ],
        ),
    ]


PERCENT_COLUMNS = {"confidence"}
WIDE_COLUMNS = {
    "text",
    "message",
    "justification",
    "excerpt",
    "action",
    "reason",
    "comment",
    "value_text",
}


def render_xlsx(report: Report) -> bytes:
    t = report.labels
    c = t["columns"]
    workbook = Workbook()
    summary = workbook.active
    summary.title = t["sheets"]["summary"]
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    summary.append([t["title"]])
    summary["A1"].font = Font(bold=True, size=14)
    if report.exploratory:
        summary.append([t["exploratory"]["title"], report.exploratory_text])
        for cell in summary[summary.max_row]:
            cell.font = Font(bold=True, color=WARNING)
            cell.fill = PatternFill("solid", fgColor="FEF3F2")
            cell.alignment = wrap
    summary.append([])
    for name, value in report.meta_rows():
        summary.append([name, value])
        summary.cell(summary.max_row, 1).font = bold
    summary.append([])
    summary.append([t["sections"]["limits"]])
    summary.cell(summary.max_row, 1).font = bold
    for line in report.limits:
        summary.append([None, line])
        summary.cell(summary.max_row, 2).alignment = wrap
    summary.append([])
    for layer in ("a", "b", "c"):
        summary.append([None, t["layers"][layer]])
        summary.cell(summary.max_row, 2).alignment = wrap
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 100

    keys_by_header = {value: key for key, value in c.items()}
    for title, headers, rows in _sheets(report):
        sheet = workbook.create_sheet(title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor="F2F4F7")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            key = keys_by_header.get(header, "")
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = (
                70
                if key in WIDE_COLUMNS
                else 28
                if key
                in {
                    "pillar",
                    "criterion",
                    "variable",
                    "scope",
                    "origin",
                    "subject",
                    "document",
                    "section",
                }
                else 14
            )
            for cell in sheet[letter][1:]:
                if key in PERCENT_COLUMNS and isinstance(cell.value, float):
                    cell.number_format = "0%"
                if key in WIDE_COLUMNS:
                    cell.alignment = wrap

    workbook.properties.title = f"{t['title']} · {report.workspace}"
    workbook.properties.creator = report.exported_by
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# Une cellule commençant par ces signes serait lue comme une formule par le
# tableur (même règle que l'export du journal d'audit).
FORMULA_PREFIX = re.compile(r"^[=+\-@\t\r]")


def _csv_cell(value: Cell, number: Callable[[float], str]) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return number(value)
    text = str(value)
    return f"'{text}" if FORMULA_PREFIX.match(text) else text


def render_csv(report: Report) -> bytes:
    """Tableau unique : une ligne par élément, la couche en première colonne
    (US-22.3). Point-virgule et BOM, comme l'export du journal d'audit."""
    t = report.labels
    c = t["columns"]
    layer = t["layer_codes"]
    rows: list[list[Cell]] = []
    if report.exploratory:
        rows.append(
            [
                layer["run"],
                t["exploratory"]["title"],
                None,
                None,
                None,
                None,
                report.exploratory_text,
                None,
            ]
        )
    rows += [
        [layer["run"], name, value, None, None, None, None, None]
        for name, value in report.meta_rows()
    ]
    rows += [
        [layer["limits"], None, None, None, None, None, line, None]
        for line in report.limits
    ]
    for pillar in report.pillars:
        for score in pillar.scores:
            rows.append(
                [
                    layer["a"],
                    c["pillar"],
                    pillar.name,
                    score.score,
                    t["evolution"].get(score.evolution or "", None),
                    score.confidence,
                    f"{c['version']} {score.version}",
                    report.refs(pillar.notes),
                ]
            )
        for variable in pillar.variables:
            rows.append(
                [
                    layer["a"],
                    c["variable"],
                    f"{pillar.name} · {variable.label}",
                    variable.value,
                    t["state"].get(variable.state, variable.state),
                    variable.confidence,
                    f"{c['version']} {variable.version}",
                    report.refs(variable.notes),
                ]
            )
    for criterion in report.criteria:
        rows.append(
            [
                layer["b"],
                c["criterion"],
                f"{criterion.code} · {criterion.name}",
                criterion.note,
                t["status"].get(criterion.status, criterion.status),
                None,
                f"{c['sub_questions']} {criterion.documented}/{criterion.total}"
                + (f" · {criterion.justification}" if criterion.justification else ""),
                report.refs(criterion.notes),
            ]
        )
    for alert in report.alerts:
        rows.append(
            [
                layer["c"],
                t["severity"].get(alert.severity, alert.severity),
                f"{alert.rule} · {alert.scope}",
                None,
                t["alert_status"].get(alert.status, alert.status),
                None,
                alert.message + (f" · {alert.comment}" if alert.comment else ""),
                report.refs(alert.notes),
            ]
        )
    for item in report.synthesis:
        rows.append(
            [
                layer["synthesis"],
                t["synthesis_sections"].get(item.section, item.section),
                item.origin,
                item.position,
                None,
                None,
                item.text,
                report.refs(item.notes),
            ]
        )
    for block in paragraphs(report.conclusion):
        rows.append(
            [layer["synthesis"], t["conclusion"], None, None, None, None, block, None]
        )
    persona = t["personas"].get(report.persona, report.persona)
    for action in report.actions:
        rows.append(
            [
                layer["actions"],
                persona,
                action.origin,
                action.position,
                None,
                None,
                f"{action.text} {t['because'].format(reason=action.justification)}",
                report.refs(action.notes),
            ]
        )
    for note in report.notes:
        rows.append(
            [
                layer["notes"],
                f"[{note.number}]",
                _source(report, note),
                None,
                None,
                None,
                note.excerpt,
                note.subject,
            ]
        )

    def number(value: float) -> str:
        text = f"{value:g}"
        return text.replace(".", ",") if report.locale == "fr" else text

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(
        [
            c["layer"],
            c["section"],
            c["element"],
            c["value"],
            c["status"],
            c["confidence"],
            c["detail"],
            c["notes"],
        ]
    )
    for row in rows:
        writer.writerow([_csv_cell(cell, number) for cell in row])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


RENDERERS: dict[str, tuple[Callable[[Report], bytes], str, str]] = {
    "pdf": (render_pdf, "application/pdf", "pdf"),
    "docx": (
        render_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "xlsx": (
        render_xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "csv": (render_csv, "text/csv; charset=utf-8", "csv"),
}
