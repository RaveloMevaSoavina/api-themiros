"""Extraction du texte (spec §4) : PDF natif, scanné ou mixte, DOCX, XLSX.

Le texte est conservé page par page, intact (RG-4.6). Un PDF scanné passe
par l'OCR Tesseract page par page, en parallèle.
"""

import logging
import zipfile
from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, replace
from io import BytesIO
from typing import Literal
from xml.etree import ElementTree

import cv2
import numpy as np
import pymupdf
import pytesseract
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.core.errors import ApiError
from app.modules.ingestion.errors import (
    corrupted_file,
    ocr_unavailable,
    password_protected,
    unsupported_format,
)

logger = logging.getLogger(__name__)

DocumentFormat = Literal["pdf", "docx", "xlsx"]
PdfType = Literal["natif", "scan", "mixte"]

# Spec §4.2 : une page de moins de 100 caractères qui contient une image est
# une page scannée.
SCANNED_PAGE_MAX_CHARS = 100
WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass(frozen=True, slots=True)
class OcrOptions:
    languages: str = "fra+eng"
    dpi: int = 300
    page_timeout_seconds: int = 120
    parallelism: int = 4


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    page_number: int
    text: str
    section_title: str | None = None
    is_ocr: bool = False
    ocr_failed: bool = False
    char_start: int = 0

    def is_exploitable(self, min_chars: int) -> bool:
        return len(self.text.strip()) >= min_chars


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    format: DocumentFormat
    pages: list[ExtractedPage]
    pdf_type: PdfType | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)

    @property
    def total_pages(self) -> int:
        return len(self.pages)

    def exploitable_pages(self, min_chars: int) -> int:
        return sum(page.is_exploitable(min_chars) for page in self.pages)


def detect_format(content: bytes, filename: str) -> DocumentFormat:
    """Le contenu fait foi, pas l'extension déclarée."""
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile as error:
            raise corrupted_file(filename) from error
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
    raise unsupported_format(filename)


def extract_document(
    content: bytes, filename: str, ocr: OcrOptions | None = None
) -> ExtractedDocument:
    document_format = detect_format(content, filename)
    try:
        if document_format == "pdf":
            extracted = _extract_pdf(content, ocr or OcrOptions())
        elif document_format == "docx":
            extracted = _extract_docx(content)
        else:
            extracted = _extract_xlsx(content)
    except ApiError:
        raise
    except Exception as error:
        raise corrupted_file(filename) from error
    return _with_offsets(extracted)


def detect_pdf_type(document: pymupdf.Document) -> PdfType:
    """Algorithme de la spec §4.2."""
    text_chars = 0
    image_pages = 0
    for page in document:
        text = page.get_text("text").strip()
        text_chars += len(text)
        if len(text) < SCANNED_PAGE_MAX_CHARS and page.get_images():
            image_pages += 1
    if text_chars < SCANNED_PAGE_MAX_CHARS * max(1, len(document)):
        return "scan"
    if image_pages:
        return "mixte"
    return "natif"


def _extract_pdf(content: bytes, ocr: OcrOptions) -> ExtractedDocument:
    document = pymupdf.open(stream=content, filetype="pdf")
    try:
        if document.needs_pass:
            raise password_protected()
        pdf_type = detect_pdf_type(document)
        pages: list[ExtractedPage] = []
        to_ocr: list[int] = []
        for index, page in enumerate(document):
            text = page.get_text("text", sort=True).strip()
            tables = _pdf_tables(page)
            if tables:
                text = f"{text}\n\n{tables}".strip()
            needs_ocr = len(text) < SCANNED_PAGE_MAX_CHARS and (
                bool(page.get_images()) or pdf_type == "scan"
            )
            if needs_ocr:
                to_ocr.append(index)
            pages.append(ExtractedPage(page_number=index + 1, text=text))

        if to_ocr:
            # PyMuPDF n'est pas thread-safe : le rendu est séquentiel, l'OCR
            # (processus Tesseract) est parallélisé.
            images = (
                (index, _render_page(document[index], ocr.dpi)) for index in to_ocr
            )
            for index, text, failed in _ocr_pages(images, ocr):
                pages[index] = replace(
                    pages[index],
                    text=text or pages[index].text,
                    is_ocr=True,
                    ocr_failed=failed,
                )
        return ExtractedDocument("pdf", pages, pdf_type)
    finally:
        document.close()


def _pdf_tables(page: pymupdf.Page) -> str:
    try:
        found = page.find_tables()
    except Exception:  # noqa: BLE001 — une page sans structure exploitable
        return ""
    blocks: list[str] = []
    for table in found.tables:
        rows = [
            " | ".join((cell or "").replace("\n", " ").strip() for cell in row)
            for row in table.extract()
        ]
        rows = [row for row in rows if row.strip(" |")]
        if rows:
            blocks.append("\n".join(rows))
    return "\n\n".join(blocks)


def _render_page(page: pymupdf.Page, dpi: int) -> np.ndarray:
    pixmap = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY, alpha=False)
    return np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
        pixmap.height, pixmap.width
    )


def _ocr_pages(
    images: Iterator[tuple[int, np.ndarray]], options: OcrOptions
) -> Iterator[tuple[int, str, bool]]:
    with ThreadPoolExecutor(max_workers=options.parallelism) as pool:
        pending = []
        for index, image in images:
            pending.append((index, pool.submit(ocr_image, image, options)))
            # Borne la mémoire : on n'attend pas la fin du rendu pour lire.
            while len(pending) >= options.parallelism * 2:
                done_index, future = pending.pop(0)
                yield done_index, *future.result()
        for index, future in pending:
            yield index, *future.result()


def ocr_image(gray: np.ndarray, options: OcrOptions) -> tuple[str, bool]:
    """OCR d'une page (spec §4.4). Retourne (texte, échec)."""
    image = preprocess_for_ocr(gray)
    try:
        text = pytesseract.image_to_string(
            image,
            lang=options.languages,
            config="--psm 3 --oem 1",
            timeout=options.page_timeout_seconds,
        )
    except pytesseract.TesseractNotFoundError as error:
        raise ocr_unavailable() from error
    except (RuntimeError, pytesseract.TesseractError):
        # Spec §4.7 : OCR échoué → extraction partielle et alerte.
        logger.warning("OCR failed on one page", exc_info=True)
        return "", True
    return text.strip(), False


def preprocess_for_ocr(gray: np.ndarray) -> np.ndarray:
    """Niveaux de gris, binarisation d'Otsu, redressement, débruitage."""
    if gray.ndim == 3:
        gray = cv2.cvtColor(gray, cv2.COLOR_RGB2GRAY)
    binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    binary = deskew(binary)
    return cv2.medianBlur(binary, 3)


def deskew(binary: np.ndarray) -> np.ndarray:
    ink = np.column_stack(np.where(binary < 128))
    if len(ink) < 50:
        return binary
    angle = cv2.minAreaRect(ink.astype(np.float32))[-1]
    if angle > 45:
        angle -= 90
    # Au-delà, l'estimation est plus probablement fausse que la page penchée.
    if abs(angle) < 0.3 or abs(angle) > 15:
        return binary
    height, width = binary.shape[:2]
    matrix = cv2.getRotationMatrix2D((width / 2, height / 2), angle, 1.0)
    return cv2.warpAffine(
        binary,
        matrix,
        (width, height),
        flags=cv2.INTER_CUBIC,
        borderMode=cv2.BORDER_REPLICATE,
    )


def _extract_docx(content: bytes) -> ExtractedDocument:
    """Paragraphes, titres, listes et tableaux dans l'ordre du document, puis
    notes de bas de page (spec §4.5). Un DOCX n'a pas de pages : une seule."""
    document = DocxDocument(BytesIO(content))
    blocks: list[str] = []
    first_heading: str | None = None
    for child in document.element.body.iterchildren():
        if child.tag == f"{WORD_NAMESPACE}p":
            paragraph = Paragraph(child, document)
            value = paragraph.text.strip()
            if not value:
                continue
            style = (paragraph.style.name if paragraph.style else "") or ""
            if style.startswith(("Heading", "Titre", "Title")):
                first_heading = first_heading or value
            elif style.startswith(("List", "Liste")):
                value = f"- {value}"
            blocks.append(value)
        elif child.tag == f"{WORD_NAMESPACE}tbl":
            table = Table(child, document)
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            rows = [row for row in rows if row.strip(" |")]
            if rows:
                blocks.append("\n".join(rows))
    footnotes = _docx_footnotes(content)
    if footnotes:
        blocks.append("\n".join(footnotes))
    return ExtractedDocument(
        "docx",
        [
            ExtractedPage(
                page_number=1, text="\n\n".join(blocks), section_title=first_heading
            )
        ],
    )


def _docx_footnotes(content: bytes) -> list[str]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if "word/footnotes.xml" not in archive.namelist():
            return []
        root = ElementTree.fromstring(archive.read("word/footnotes.xml"))
    notes: list[str] = []
    for footnote in root.iter(f"{WORD_NAMESPACE}footnote"):
        if int(footnote.attrib.get(f"{WORD_NAMESPACE}id", "-1")) < 1:
            continue
        text = "".join(node.text or "" for node in footnote.iter(f"{WORD_NAMESPACE}t"))
        if text.strip():
            notes.append(text.strip())
    return notes


def _extract_xlsx(content: bytes) -> ExtractedDocument:
    """Une page par feuille, cellules repérées par leur coordonnée, valeurs
    calculées des formules (spec §4.6)."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        pages: list[ExtractedPage] = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines: list[str] = []
            for row in sheet.iter_rows():
                values = [
                    f"{cell.coordinate}={cell.value}"
                    for cell in row
                    if cell.value is not None and str(cell.value).strip()
                ]
                if values:
                    lines.append(" | ".join(values))
            pages.append(
                ExtractedPage(
                    page_number=index, text="\n".join(lines), section_title=sheet.title
                )
            )
        return ExtractedDocument("xlsx", pages)
    finally:
        workbook.close()


def _with_offsets(document: ExtractedDocument) -> ExtractedDocument:
    """Position de chaque page dans le texte complet (pages jointes par une
    ligne vide), pour les `char_start` / `char_end` des segments."""
    offset = 0
    pages: list[ExtractedPage] = []
    for page in document.pages:
        pages.append(replace(page, char_start=offset))
        offset += len(page.text) + 2
    return replace(document, pages=pages)
